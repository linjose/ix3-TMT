import tempfile
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from pptx import Presentation
from pptx.util import Inches

from .models import Collection, Deck, Slide, Star
from .services.text_extract import extract_native_slides, extract_keywords

User = get_user_model()


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class TMTModelAndViewTests(TestCase):
    def setUp(self):
        self.alice = User.objects.create_user(username="alice", password="pw-strong-123")
        self.bob = User.objects.create_user(username="bob", password="pw-strong-123")
        self.deck = Deck.objects.create(
            owner=self.alice, title="AI Strategy", original_filename="ai.pptx",
            original_file=SimpleUploadedFile("ai.pptx", b"dummy"), status=Deck.Status.READY,
            visibility=Deck.Visibility.INTERNAL, page_count=1,
        )
        self.slide = Slide.objects.create(
            deck=self.deck, page_number=1, title="AI Agent Architecture",
            native_text="AI Agent Architecture", search_text="AI Agent Architecture healthcare",
            image=SimpleUploadedFile("1.png", b"not-image"),
            thumbnail=SimpleUploadedFile("1.webp", b"not-image"),
        )

    def test_internal_slide_visible_to_authenticated_user(self):
        self.assertTrue(Slide.objects.visible_to(self.bob).filter(pk=self.slide.pk).exists())

    def test_private_deck_not_visible_to_other_user(self):
        self.deck.visibility = Deck.Visibility.PRIVATE
        self.deck.save(update_fields=["visibility"])
        self.assertFalse(Slide.objects.visible_to(self.bob).filter(pk=self.slide.pk).exists())

    def test_star_toggle_updates_counter(self):
        self.client.login(username="bob", password="pw-strong-123")
        response = self.client.post(reverse("tmt:toggle_star", args=[self.slide.pk]))
        self.assertEqual(response.status_code, 200)
        self.slide.refresh_from_db()
        self.assertEqual(self.slide.star_count, 1)
        self.assertTrue(Star.objects.filter(user=self.bob, slide=self.slide).exists())

    def test_collection_can_hold_same_slide_independently(self):
        c1 = Collection.objects.create(user=self.bob, name="AI")
        c2 = Collection.objects.create(user=self.bob, name="Proposal")
        c1.slides.add(self.slide)
        c2.slides.add(self.slide)
        self.assertEqual(self.slide.collections.filter(user=self.bob).count(), 2)

    def test_search_finds_slide_text(self):
        self.client.login(username="bob", password="pw-strong-123")
        response = self.client.get(reverse("tmt:home"), {"q": "healthcare"})
        self.assertContains(response, "AI Agent Architecture")

    def test_health_endpoint(self):
        response = self.client.get(reverse("tmt:health"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["ok"], True)


class TextExtractionTests(TestCase):
    def test_extract_native_pptx_text(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "sample.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
            box.text = "顧問級 AI Agent Architecture"
            prs.save(path)
            pages = extract_native_slides(path)
            self.assertEqual(len(pages), 1)
            self.assertIn("AI Agent Architecture", pages[0].text)

    def test_keywords_support_mixed_language(self):
        tags = extract_keywords("AI Agent 架構與人工智慧 Agent platform platform", topk=5)
        lowered = [t.lower() for t in tags]
        self.assertIn("platform", lowered)
